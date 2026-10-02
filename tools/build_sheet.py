"""Build the Ýmir ops workbook from the schema and the extracted data.

    python tools/build_sheet.py <data_dir> <out.xlsx>

The data directory holds:
  ops_<Type>.yaml       rows parsed from the vault's Operations/ notes (tools/parse_ops.py)
  ops_links.yaml        links produced by the parser
  curated_*.yaml        hand-extracted rows: {<Type>: [rows], links: [[from, relation, to, source]]}

Rows in curated files with the same id as a parsed row are merged into it.
The build stops with a list of problems if any id is duplicated, a link points
at a missing row, a relation is unknown, or a row has a field the schema lacks.
"""
import re
import sys
from datetime import date
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
ISO_DATE = re.compile(r"\d{4}-\d{2}-\d{2}")
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10)
CONFLICT_FILL = PatternFill("solid", fgColor="FFF2CC")
WIDE = {"name": 48, "text": 60, "why": 48, "options": 40, "covers": 44, "scope": 44, "remit": 44,
        "must_be_true": 44, "relationship": 40, "note": 40, "conflict": 40, "quote": 44, "source": 36}
COMMON_VALUES = {"basis": ["notes", "proposal", "mixed"], "origin": ["extracted", "human"]}


def load_schema():
    return yaml.safe_load((ROOT / "schema" / "types.yaml").read_text())


def load_data(data_dir, schema):
    types = {t["name"]: t for t in schema["types"]}
    rows = {name: {} for name in types}
    links, problems = [], []

    def add(type_name, row, origin_file):
        if type_name not in types:
            problems.append(f"{origin_file}: unknown type {type_name}")
            return
        table = rows[type_name]
        if row["id"] in table:
            if origin_file.startswith("ops_") or table[row["id"]].get("_file", "").startswith("curated"):
                problems.append(f"{origin_file}: duplicate id {row['id']}")
                return
            table[row["id"]].update(row)
        else:
            table[row["id"]] = dict(row, _file=origin_file)

    for path in sorted(Path(data_dir).glob("ops_*.yaml")):
        content = yaml.safe_load(path.read_text())
        if path.stem == "ops_links":
            links += [[l["from"], l["relation"], l["to"], l.get("source", "")] for l in content]
        else:
            for row in content:
                add(path.stem[4:], row, path.name)
    for path in sorted(Path(data_dir).glob("curated_*.yaml")):
        content = yaml.safe_load(path.read_text())
        for key, value in content.items():
            if key == "links":
                links += value
            else:
                for row in value:
                    add(key, row, path.name)
    return rows, links, problems


def finish_rows(rows, schema):
    """Fill derived columns and flag fields the schema doesn't know."""
    problems = []
    common = [c["name"] for c in schema["common"]]
    for t in schema["types"]:
        allowed = {"id", "name", *common, *(f["name"] for f in t["fields"])}
        for row in rows[t["name"]].values():
            origin_file = row.pop("_file", "")
            extra = set(row) - allowed
            if extra:
                problems.append(f"{t['name']} {row['id']}: unknown field(s) {sorted(extra)}")
            row.setdefault("origin", "extracted")
            for col in ("amount_isk", "amount"):
                value = str(row.get(col, ""))
                if value.replace(".", "", 1).isdigit():
                    row[col] = float(value) if "." in value else int(value)
            # Parsed rows whose content also appears in an ordinary note are "mixed":
            # the item is in the notes, the owner/date/rank came from the Operations proposal.
            if origin_file.startswith("ops_") and row.get("basis") == "proposal":
                notes = [s for s in row.get("source", "").split("; ") if s and not s.startswith("Operations")]
                if notes and t["name"] in ("Task", "Decision"):
                    row["basis"] = "mixed"
    return problems


def resolve_links(rows, links, schema):
    """Resolve ~Type:text references and check every link."""
    problems = []
    id_type = {rid: tname for tname, table in rows.items() for rid in table}
    relations = {r["name"]: r for r in schema["relations"]}
    resolved = []
    for frm, rel, to, src in links:
        ends = []
        for ref in (frm, to):
            if ref.startswith("~"):
                tname, text = ref[1:].split(":", 1)
                hits = [rid for rid, r in rows[tname].items() if text.lower() in r["name"].lower()]
                if len(hits) != 1:
                    problems.append(f"link {ref}: {len(hits)} matches")
                    ref = None
                else:
                    ref = hits[0]
            elif ref not in id_type:
                problems.append(f"link {frm} {rel} {to}: no row {ref}")
                ref = None
            ends.append(ref)
        if None in ends:
            continue
        spec = relations.get(rel)
        if not spec:
            problems.append(f"link {frm} {rel} {to}: unknown relation")
            continue
        for end, want in zip(ends, (spec["from"], spec["to"])):
            if want != "*" and id_type[end] != want:
                problems.append(f"link {frm} {rel} {to}: {end} is a {id_type[end]}, relation wants {want}")
        resolved.append([ends[0], id_type[ends[0]], rel, ends[1], id_type[ends[1]], src])
    seen, unique = set(), []
    for link in resolved:
        key = (link[0], link[2], link[3])
        if key not in seen:
            seen.add(key)
            unique.append(link)
    return unique, problems


def style_header(ws, columns):
    for i, col in enumerate(columns, 1):
        cell = ws.cell(row=1, column=i, value=col)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = WIDE.get(col, 16 if col != "id" else 22)
    ws.freeze_panes = "C2"
    ws.row_dimensions[1].height = 22


def add_dropdown(ws, col_index, values, n_rows):
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
    dv.error, dv.showErrorMessage = "Pick a value from the list, or add it to schema/types.yaml", False
    letter = get_column_letter(col_index)
    dv.add(f"{letter}2:{letter}{max(n_rows + 200, 300)}")
    ws.add_data_validation(dv)


def write_type_tab(wb, t, table, schema):
    ws = wb.create_sheet(t["name"])
    fields = [f["name"] for f in t["fields"]]
    common = [c["name"] for c in schema["common"]]
    columns = ["id", "name", *fields, *common, "links"]
    style_header(ws, columns)
    links_col = get_column_letter(len(columns))
    for r, row in enumerate(sorted(table.values(), key=sort_key), 2):
        for c, col in enumerate(columns[:-1], 1):
            value = row.get(col, "")
            is_date = isinstance(value, str) and ISO_DATE.fullmatch(value)
            if is_date:
                value = date.fromisoformat(value)
            elif isinstance(value, str) and value and value[0] in "=+-@":
                value = "'" + value
            cell = ws.cell(row=r, column=c, value=value if value != "" else None)
            if is_date:
                cell.number_format = "yyyy-mm-dd"
            cell.font = BODY_FONT
            cell.alignment = Alignment(wrap_text=col in WIDE, vertical="top")
            if col == "conflict" and value:
                cell.fill = CONFLICT_FILL
    # One formula fills the whole column, so rows added later are counted too.
    ws.cell(row=2, column=len(columns),
            value='=ARRAYFORMULA(IF(A2:A="","",COUNTIF(Links!A:A,A2:A)+COUNTIF(Links!D:D,A2:A)))').font = BODY_FONT
    for i, col in enumerate(columns, 1):
        spec = next((f for f in t["fields"] if f["name"] == col), None)
        values = (spec or {}).get("values") or (t["status"] if col == "status" else COMMON_VALUES.get(col))
        if values:
            add_dropdown(ws, i, [str(v) for v in values], len(table))
    for col in ("amount_isk", "amount"):
        if col in columns:
            letter = get_column_letter(columns.index(col) + 1)
            for cell in ws[letter][1:]:
                cell.number_format = "#,##0;(#,##0);0"
    ws.auto_filter.ref = f"A1:{links_col}{len(table) + 1}"
    return ws


def sort_key(row):
    rid = row["id"]
    parts = rid.replace("D", "D-", 1).split("-") if rid[0] == "D" and rid[1:].isdigit() else rid.split("-")
    return [p.zfill(6) if p.replace(".", "").isdigit() else p for p in parts]


def write_links(wb, links):
    ws = wb.create_sheet("Links")
    columns = ["from", "from_type", "relation", "to", "to_type", "source", "from_name", "to_name", "origin"]
    style_header(ws, columns)
    for c, w in zip("ABCDEFGHI", (18, 12, 14, 18, 12, 36, 44, 44, 10)):
        ws.column_dimensions[c].width = w
    for r, (frm, ft, rel, to, tt, src) in enumerate(links, 2):
        for c, v in enumerate([frm, ft, rel, to, tt, src, None, None, "extracted"], 1):
            if v is not None:
                ws.cell(row=r, column=c, value=v).font = BODY_FONT
    for col, key in (("G", "A"), ("H", "D")):
        ws[f"{col}2"] = (f'=ARRAYFORMULA(IF({key}2:{key}="","",'
                         f'IFERROR(VLOOKUP({key}2:{key},_Index!A:B,2,FALSE),"(missing)")))')
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{len(links) + 1}"


def write_meta(wb, schema):
    ws = wb.create_sheet("_Types")
    columns = ["type", "group", "id prefix", "description", "fields", "status values"]
    style_header(ws, columns)
    for c, w in zip("ABCDEF", (14, 12, 10, 60, 50, 36)):
        ws.column_dimensions[c].width = w
    for r, t in enumerate(schema["types"], 2):
        for c, v in enumerate([t["name"], t["group"], t["prefix"], t["description"],
                               ", ".join(f["name"] for f in t["fields"]), ", ".join(t["status"])], 1):
            ws.cell(row=r, column=c, value=v).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"

    ws = wb.create_sheet("_Relations")
    style_header(ws, ["relation", "from", "to", "meaning"])
    for c, w in zip("ABCD", (14, 14, 14, 60)):
        ws.column_dimensions[c].width = w
    for r, rel in enumerate(schema["relations"], 2):
        for c, v in enumerate([rel["name"], rel["from"], rel["to"], rel["meaning"]], 1):
            ws.cell(row=r, column=c, value=v)
    ws.freeze_panes = "A2"


def column_letter(t, schema, name):
    common = [c["name"] for c in schema["common"]]
    columns = ["id", "name", *(f["name"] for f in t["fields"]), *common, "links"]
    return get_column_letter(columns.index(name) + 1)


def write_review(wb, schema):
    """Live list of every row with a conflict, across all type tabs."""
    ws = wb.create_sheet("Review", 1)
    style_header(ws, ["id", "name", "conflict", "source"])
    for c, w in zip("ABCD", (18, 48, 60, 40)):
        ws.column_dimensions[c].width = w
    parts = []
    for t in schema["types"]:
        n, cf, src = t["name"], column_letter(t, schema, "conflict"), column_letter(t, schema, "source")
        parts.append(f"{{{n}!A2:B,{n}!{cf}2:{cf},{n}!{src}2:{src}}}")
    ws["A2"] = "=QUERY({" + ";".join(parts) + "},\"select * where Col3 is not null and Col3 <> ''\",0)"
    ws.freeze_panes = "A2"


def write_index(wb, schema):
    """Every id and name in one place, for the name lookups on the Links tab."""
    ws = wb.create_sheet("_Index")
    style_header(ws, ["id", "name"])
    stacked = ";".join(f"{t['name']}!A2:B" for t in schema["types"])
    ws["A2"] = f'=QUERY({{{stacked}}},"select * where Col1 is not null",0)'


def write_start(wb, schema):
    ws = wb.active
    ws.title = "Start here"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 70
    lines = [
        ("Ýmir ops data", None, None),
        ("Built from the Obsidian vault (skarfur/ymir_vault). Each tab below is one type of thing; "
         "every row has an id, the note it came from, and a quote where useful. Links between rows are in the Links tab.", None, None),
        ("", None, None),
        ("Column", None, "Meaning"),
        ("basis", None, "notes = stated in the notes; proposal = only in the Operations/ synthesis; "
                        "mixed = the item is in the notes, its owner/date/rank were proposed"),
        ("conflict", None, "Two notes disagree. All of these are listed on the Review tab"),
        ("origin", None, "extracted = written by the extraction. Change it to human when you edit a row, "
                         "and later extractions will leave that row alone"),
        ("links", None, "How many links touch this row (formula)"),
        ("", None, None),
        ("Tab", "Rows", "What it holds"),
    ]
    for t in schema["types"]:
        lines.append((t["name"], f"=COUNTA('{t['name']}'!A:A)-1", t["description"]))
    lines += [("Links", "=COUNTA(Links!A:A)-1", "Every relationship between rows: from, relation, to, and the note it came from"),
              ("Review", "=COUNTA(Review!A:A)-1", "Rows where the notes contradict each other"),
              ("_Types / _Relations", None, "The schema. To add a type or field, change schema/types.yaml in skarfur/ymir_ops and rebuild"),
              ("Broken links", '=COUNTIF(Links!G:H,"(missing)")', "Links whose from or to id doesn't exist in any tab. Should be 0; filter the Links tab for (missing) to find them"),
              ("_Index", "=COUNTA(_Index!A:A)-1", "Every id and name across all tabs, used by the Links name columns. Don't edit")]
    for r, (a, b, c) in enumerate(lines, 1):
        ws.cell(row=r, column=1, value=a)
        if b:
            ws.cell(row=r, column=2, value=b)
        if c:
            ws.cell(row=r, column=3, value=c).alignment = Alignment(wrap_text=True, vertical="top")
    ws["A1"].font = Font(name="Arial", size=14, bold=True)
    ws.merge_cells("A2:C2")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 45
    for r in (4, 10):
        for c in range(1, 4):
            ws.cell(row=r, column=c).font = Font(name="Arial", bold=True)


def main(data_dir, out):
    schema = load_schema()
    rows, links, problems = load_data(data_dir, schema)
    problems += finish_rows(rows, schema)
    links, link_problems = resolve_links(rows, links, schema)
    problems += link_problems
    if problems:
        print("Problems:\n  " + "\n  ".join(problems))
        sys.exit(1)

    wb = Workbook()
    write_start(wb, schema)
    for t in schema["types"]:
        write_type_tab(wb, t, rows[t["name"]], schema)
    write_links(wb, links)
    write_review(wb, schema)
    write_meta(wb, schema)
    write_index(wb, schema)
    wb.save(out)
    for t in schema["types"]:
        print(f"{t['name']:<13}{len(rows[t['name']]):>4}")
    print(f"{'Links':<13}{len(links):>4}")
    print(f"{'Conflicts':<13}{sum(1 for tb in rows.values() for r in tb.values() if r.get('conflict')):>4}")


if __name__ == "__main__":
    main(*sys.argv[1:3])
